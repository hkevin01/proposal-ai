"""
Export Utility
Exports analytics and proposals to PDF and DOCX formats.
"""
from fpdf import FPDF
from docx import Document
from typing import List, Dict, Any
import yaml
import os

EXPORT_CONFIG_PATH = "config/export_settings.yaml"


def load_export_settings() -> dict:
    """Load export settings from YAML config."""
    try:
        with open(EXPORT_CONFIG_PATH, "r") as f:
            return yaml.safe_load(f)
    except Exception:
        return {}


def export_to_pdf(data: List[Dict[str, Any]], filename: str = None) -> None:
    """Export data to PDF with config options."""
    settings = load_export_settings().get("pdf", {})
    font = settings.get("default_font", "Arial")
    size = settings.get("default_size", 12)
    output_dir = settings.get("output_dir", "exports/pdf")
    if not filename:
        filename = os.path.join(output_dir, "analytics_report.pdf")
    os.makedirs(output_dir, exist_ok=True)
    try:
        pdf = FPDF()
        pdf.add_page()
        pdf.set_font(font, size=size)
        for item in data:
            pdf.cell(200, 10, txt=str(item), ln=True)
        pdf.output(filename)
    except Exception as e:
        print(f"Error exporting to PDF: {e}")


def export_to_docx(data: List[Dict[str, Any]], filename: str = None) -> None:
    """Export data to DOCX with config options."""
    settings = load_export_settings().get("docx", {})
    output_dir = settings.get("output_dir", "exports/docx")
    if not filename:
        filename = os.path.join(output_dir, "analytics_report.docx")
    os.makedirs(output_dir, exist_ok=True)
    try:
        doc = Document()
        for item in data:
            doc.add_paragraph(str(item))
        doc.save(filename)
    except Exception as e:
        print(f"Error exporting to DOCX: {e}")
